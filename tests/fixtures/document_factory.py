"""Small in-memory PDF and DOCX factories for tests."""

from io import BytesIO

import pymupdf
from docx import Document


def make_pdf(*pages: str) -> bytes:
    document = pymupdf.open()
    try:
        for content in pages:
            page = document.new_page()
            if content:
                page.insert_text((72, 72), content)
        return document.tobytes()
    finally:
        document.close()


def make_docx(
    paragraphs: tuple[str, ...] = (),
    table_rows: tuple[tuple[str, ...], ...] = (),
) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        column_count = max(len(row) for row in table_rows)
        table = document.add_table(rows=len(table_rows), cols=column_count)
        for row_index, values in enumerate(table_rows):
            for column_index, value in enumerate(values):
                table.cell(row_index, column_index).text = value

    output = BytesIO()
    document.save(output)
    return output.getvalue()
